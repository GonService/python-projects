"""Turns a sales spreadsheet (CSV or XLSX) into a summary report."""

from __future__ import annotations

import csv
import io
from collections import defaultdict
from dataclasses import dataclass, field
from datetime import date, datetime
from decimal import Decimal, InvalidOperation

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font

REQUIRED = ("date", "product", "category", "quantity", "unit_price")
OPTIONAL = ("seller",)
ALIASES = {
    "data": "date",
    "produto": "product",
    "categoria": "category",
    "quantidade": "quantity",
    "qtd": "quantity",
    "preco_unitario": "unit_price",
    "preço_unitário": "unit_price",
    "valor_unitario": "unit_price",
    "vendedor": "seller",
}


class ReportError(ValueError):
    """Problem with the uploaded file, with a message the user can act on."""


@dataclass
class Sale:
    day: date
    product: str
    category: str
    quantity: int
    unit_price: Decimal
    seller: str = ""

    @property
    def total(self) -> Decimal:
        return self.unit_price * self.quantity


@dataclass
class Summary:
    revenue: Decimal = Decimal("0")
    orders: int = 0
    units: int = 0
    first_day: date | None = None
    last_day: date | None = None
    by_month: dict[str, Decimal] = field(default_factory=dict)
    by_category: list[tuple[str, Decimal]] = field(default_factory=list)
    top_products: list[tuple[str, int, Decimal]] = field(default_factory=list)
    by_seller: list[tuple[str, Decimal]] = field(default_factory=list)
    skipped_rows: list[str] = field(default_factory=list)

    @property
    def average_ticket(self) -> Decimal:
        return (self.revenue / self.orders).quantize(Decimal("0.01")) if self.orders else Decimal("0")

    def as_dict(self) -> dict:
        money = lambda v: str(v.quantize(Decimal("0.01")))  # noqa: E731
        return {
            "revenue": money(self.revenue),
            "orders": self.orders,
            "units": self.units,
            "average_ticket": money(self.average_ticket),
            "period": {
                "from": self.first_day.isoformat() if self.first_day else None,
                "to": self.last_day.isoformat() if self.last_day else None,
            },
            "by_month": {k: money(v) for k, v in self.by_month.items()},
            "by_category": [{"category": c, "revenue": money(v)} for c, v in self.by_category],
            "top_products": [{"product": p, "units": u, "revenue": money(v)} for p, u, v in self.top_products],
            "by_seller": [{"seller": s, "revenue": money(v)} for s, v in self.by_seller],
            "skipped_rows": self.skipped_rows,
        }


def _normalize_header(name: object) -> str:
    key = str(name or "").strip().lower().replace(" ", "_")
    return ALIASES.get(key, key)


def _parse_date(value: object) -> date:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = str(value).strip()
    for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y", "%m/%d/%Y"):
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            continue
    raise ValueError(f"invalid date '{text}'")


def _parse_decimal(value: object) -> Decimal:
    if isinstance(value, int | float | Decimal):
        return Decimal(str(value))
    text = str(value).strip().replace("R$", "").replace(" ", "")
    if "," in text and "." in text:
        # 1.234,56 (pt-BR) or 1,234.56 (en): the last separator is the decimal one
        text = text.replace(".", "").replace(",", ".") if text.rfind(",") > text.rfind(".") else text.replace(",", "")
    elif "," in text:
        text = text.replace(",", ".")
    try:
        return Decimal(text)
    except InvalidOperation as exc:
        raise ValueError(f"invalid number '{value}'") from exc


def _rows_from_csv(data: bytes) -> list[list[object]]:
    text = data.decode("utf-8-sig", errors="replace")
    try:
        dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t")
    except csv.Error:
        dialect = csv.excel
    return [row for row in csv.reader(io.StringIO(text), dialect) if any(cell.strip() for cell in row)]


def _rows_from_xlsx(data: bytes) -> list[list[object]]:
    try:
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception as exc:  # openpyxl raises several types for corrupt files
        raise ReportError("Could not open the spreadsheet. Save it as .xlsx or .csv and try again.") from exc
    ws = wb.worksheets[0]
    return [list(r) for r in ws.iter_rows(values_only=True) if any(c not in (None, "") for c in r)]


def read_sales(filename: str, data: bytes) -> tuple[list[Sale], list[str]]:
    """Parses the file. Returns valid sales and a list of skipped rows with the reason."""
    name = filename.lower()
    if name.endswith(".csv"):
        rows = _rows_from_csv(data)
    elif name.endswith(".xlsx"):
        rows = _rows_from_xlsx(data)
    else:
        raise ReportError("Upload a .csv or .xlsx file.")
    if not rows:
        raise ReportError("The file is empty.")

    header = [_normalize_header(h) for h in rows[0]]
    missing = [c for c in REQUIRED if c not in header]
    if missing:
        raise ReportError("Missing columns: " + ", ".join(missing) + ". Expected: " + ", ".join(REQUIRED) + ".")
    idx = {name: header.index(name) for name in REQUIRED + OPTIONAL if name in header}

    sales: list[Sale] = []
    skipped: list[str] = []
    for line, row in enumerate(rows[1:], start=2):
        try:

            def cell(c: str, row: list[object] = row) -> object:
                return row[idx[c]] if idx[c] < len(row) else None

            qty = _parse_decimal(cell("quantity"))
            if qty != qty.to_integral_value() or qty <= 0:
                raise ValueError("quantity must be a positive whole number")
            price = _parse_decimal(cell("unit_price"))
            if price < 0:
                raise ValueError("unit_price cannot be negative")
            product = str(cell("product") or "").strip()
            if not product:
                raise ValueError("product is empty")
            sales.append(
                Sale(
                    day=_parse_date(cell("date")),
                    product=product,
                    category=str(cell("category") or "Uncategorized").strip() or "Uncategorized",
                    quantity=int(qty),
                    unit_price=price,
                    seller=str(cell("seller") or "").strip() if "seller" in idx else "",
                )
            )
        except (ValueError, TypeError) as exc:
            skipped.append(f"Row {line}: {exc}")
    if not sales:
        raise ReportError("No valid rows found. " + (skipped[0] if skipped else ""))
    return sales, skipped


def summarize(sales: list[Sale], top: int = 5) -> Summary:
    s = Summary()
    month: dict[str, Decimal] = defaultdict(Decimal)
    category: dict[str, Decimal] = defaultdict(Decimal)
    seller: dict[str, Decimal] = defaultdict(Decimal)
    product_units: dict[str, int] = defaultdict(int)
    product_rev: dict[str, Decimal] = defaultdict(Decimal)

    for sale in sales:
        s.revenue += sale.total
        s.orders += 1
        s.units += sale.quantity
        s.first_day = min(s.first_day or sale.day, sale.day)
        s.last_day = max(s.last_day or sale.day, sale.day)
        month[sale.day.strftime("%Y-%m")] += sale.total
        category[sale.category] += sale.total
        product_units[sale.product] += sale.quantity
        product_rev[sale.product] += sale.total
        if sale.seller:
            seller[sale.seller] += sale.total

    s.by_month = dict(sorted(month.items()))
    s.by_category = sorted(category.items(), key=lambda kv: kv[1], reverse=True)
    s.top_products = sorted(
        ((p, product_units[p], product_rev[p]) for p in product_rev), key=lambda t: t[2], reverse=True
    )[:top]
    s.by_seller = sorted(seller.items(), key=lambda kv: kv[1], reverse=True)
    return s


def to_xlsx(summary: Summary) -> bytes:
    """Summary as a spreadsheet with one sheet per breakdown."""
    wb = Workbook()
    bold = Font(bold=True)

    def sheet(title: str, header: list[str], rows: list[list[object]], first: bool = False):
        ws = wb.active if first else wb.create_sheet()
        ws.title = title
        ws.append(header)
        for cell in ws[1]:
            cell.font = bold
        for r in rows:
            ws.append(r)
        for col in ws.columns:
            ws.column_dimensions[col[0].column_letter].width = max(12, *(len(str(c.value or "")) + 2 for c in col))

    sheet(
        "Summary",
        ["Metric", "Value"],
        [
            ["Revenue", float(summary.revenue)],
            ["Orders", summary.orders],
            ["Units", summary.units],
            ["Average ticket", float(summary.average_ticket)],
            ["From", summary.first_day.isoformat() if summary.first_day else ""],
            ["To", summary.last_day.isoformat() if summary.last_day else ""],
        ],
        first=True,
    )
    sheet("By month", ["Month", "Revenue"], [[m, float(v)] for m, v in summary.by_month.items()])
    sheet("By category", ["Category", "Revenue"], [[c, float(v)] for c, v in summary.by_category])
    sheet("Top products", ["Product", "Units", "Revenue"], [[p, u, float(v)] for p, u, v in summary.top_products])
    if summary.by_seller:
        sheet("By seller", ["Seller", "Revenue"], [[n, float(v)] for n, v in summary.by_seller])

    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()
