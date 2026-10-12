import io
from decimal import Decimal

import pytest
from openpyxl import Workbook, load_workbook

from app.report import ReportError, read_sales, summarize, to_xlsx

CSV = b"""date,product,category,quantity,unit_price,seller
2026-07-01,Coffee 250g,Coffee,2,34.90,Ana
2026-07-15,Grinder,Gear,1,189.00,Bruno
2026-08-02,Coffee 250g,Coffee,3,34.90,Ana
"""


def test_reads_csv_and_summarizes():
    sales, skipped = read_sales("sales.csv", CSV)
    assert skipped == []
    s = summarize(sales)
    assert s.revenue == Decimal("363.50")
    assert (s.orders, s.units) == (3, 6)
    assert s.average_ticket == Decimal("121.17")
    assert s.by_month == {"2026-07": Decimal("258.80"), "2026-08": Decimal("104.70")}
    assert s.by_category[0] == ("Gear", Decimal("189.00"))
    assert s.top_products[0] == ("Grinder", 1, Decimal("189.00"))
    assert dict(s.by_seller) == {"Ana": Decimal("174.50"), "Bruno": Decimal("189.00")}


def test_portuguese_headers_semicolons_and_brazilian_numbers():
    data = "data;produto;categoria;quantidade;preço_unitário\n01/07/2026;Café;Café;2;R$ 1.234,50\n".encode()
    sales, _ = read_sales("vendas.csv", data)
    assert sales[0].unit_price == Decimal("1234.50")
    assert sales[0].day.isoformat() == "2026-07-01"
    assert summarize(sales).revenue == Decimal("2469.00")


def test_reads_xlsx():
    wb = Workbook()
    ws = wb.active
    ws.append(["Date", "Product", "Category", "Quantity", "Unit price"])
    ws.append(["2026-09-10", "Mug", "Gear", 2, 79.9])
    buf = io.BytesIO()
    wb.save(buf)
    sales, _ = read_sales("sales.xlsx", buf.getvalue())
    assert summarize(sales).revenue == Decimal("159.8")


def test_bad_rows_are_skipped_with_a_reason():
    data = CSV + b"2026-08-03,Mug,Gear,abc,10,Ana\n2026-13-40,Mug,Gear,1,10,Ana\n2026-08-04,,Gear,1,10,Ana\n"
    sales, skipped = read_sales("sales.csv", data)
    assert len(sales) == 3
    assert [r.split(":")[0] for r in skipped] == ["Row 5", "Row 6", "Row 7"]
    assert "invalid number" in skipped[0]


def test_missing_columns_explain_what_is_expected():
    with pytest.raises(ReportError, match="Missing columns: unit_price"):
        read_sales("x.csv", b"date,product,category,quantity\n2026-01-01,a,b,1\n")


def test_rejects_other_file_types_and_empty_files():
    with pytest.raises(ReportError, match=".csv or .xlsx"):
        read_sales("x.pdf", b"%PDF")
    with pytest.raises(ReportError, match="empty"):
        read_sales("x.csv", b"")


def test_excel_export_has_one_sheet_per_breakdown():
    s = summarize(read_sales("sales.csv", CSV)[0])
    wb = load_workbook(io.BytesIO(to_xlsx(s)))
    assert wb.sheetnames == ["Summary", "By month", "By category", "Top products", "By seller"]
    assert wb["Summary"]["B2"].value == pytest.approx(363.5)
